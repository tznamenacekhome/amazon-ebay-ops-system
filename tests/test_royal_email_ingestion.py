from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "integrations"))

from royal_email_ingestion import GraphClient, IntakeError, allowed_host, body_links, validate_source
from wholesale_royal import filename_effective_date, parse_price_list


class RoyalEmailSafetyTests(unittest.TestCase):
    def test_attachment_listing_uses_graph_supported_metadata_fields(self):
        class Response:
            status_code = 200
            def json(self): return {"value": []}

        class Session:
            def __init__(self): self.params = None
            def get(self, _url, *, params, headers, timeout):
                self.params = params
                return Response()

        session = Session()
        graph = GraphClient("tenant", "client", "secret", session=session)
        graph._token = "token"
        self.assertEqual(graph.attachments("mailbox", "message"), [])
        self.assertEqual(session.params["$select"], "id,name,contentType,size,isInline")

    def test_download_domain_requires_exact_host_or_subdomain(self):
        allowed = {"portal.royalelec.com"}
        self.assertTrue(allowed_host("portal.royalelec.com", allowed))
        self.assertTrue(allowed_host("files.portal.royalelec.com", allowed))
        self.assertFalse(allowed_host("portal.royalelec.com.evil.test", allowed))

    def test_body_links_are_extracted_without_logging_or_storage(self):
        links = body_links({"content": '<a href="https://portal.royalelec.com/d?id=secret">Download</a>'})
        self.assertEqual(links, ["https://portal.royalelec.com/d?id=secret"])

    def test_binary_csv_is_rejected(self):
        with self.assertRaises(IntakeError):
            validate_source("Royal.csv", b"ORDER,TITLE\x00bad")

    def test_filename_date_is_conservative(self):
        self.assertEqual(str(filename_effective_date("Royal Price List 2026-10-03.csv")), "2026-10-03")
        self.assertEqual(str(filename_effective_date("PRICE LISTS 9082026.xlsx")), "2026-09-08")
        self.assertEqual(
            str(filename_effective_date("Royal NN 10-05.xlsx", "2026-10-04T23:41:09Z")),
            "2026-10-05",
        )
        self.assertIsNone(filename_effective_date("unrelated-2026-10-03.csv"))

    def test_price_list_sheet_requires_exact_full_list_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            accepted = Path(directory) / "Royal NN 10-05.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Price List"
            sheet.append(["ORDER", "TITLE", "SYS", "UPC / SKU", "PRICE", "QTY", "SUB"])
            sheet.append([1, "Example Game", "Switch", "012345678905", 19.99, 5, None])
            workbook.save(accepted)
            workbook.close()
            parsed = parse_price_list(accepted, filename_reference_date="2026-10-04T23:41:09Z")

            rejected = Path(directory) / "Royal Sale 10-05.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Price List"
            sheet.append(["ITEM", "DESCRIPTION", "PRICE"])
            workbook.save(rejected)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "full-list sheet and headers"):
                parse_price_list(rejected, filename_reference_date="2026-10-04T23:41:09Z")

        self.assertEqual(parsed["status"], "completed")
        self.assertEqual(parsed["effective_date"], "2026-10-05")
        self.assertEqual(parsed["date_source"], "supplier_filename")

    def test_current_royal_layout_accepts_order_total_before_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Royal NN Price_List_10-07.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "PRICE LIST"
            sheet.append(["ORDER TOTAL:", None, None, None, None, None, "=SUM(G3:G4)"])
            sheet.append(["ORDER", "TITLE", "SYS", "UPC/SKU", "PRICE", "QTY", "SUB"])
            sheet.append([None, "Example Game", "Switch", "012345678905", 19.99, "144+", '=IF(A3="","",A3*E3)'])
            workbook.save(path)
            workbook.close()

            parsed = parse_price_list(path, filename_reference_date="2026-10-06T23:36:54Z")

        self.assertEqual(parsed["status"], "completed")
        self.assertEqual(parsed["effective_date"], "2026-10-07")
        self.assertEqual(parsed["date_source"], "supplier_filename")
        self.assertEqual(parsed["summary"]["rows_imported"], 1)
        self.assertEqual(parsed["products"][0]["normalized_identifier"], "012345678905")

    def test_csv_uses_existing_royal_normalization(self):
        content = "ORDER,TITLE,SYS,UPC / SKU,PRICE,QTY,SUB\n1,Example Game,Switch,012345678905,19.99,5,\n"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Royal Price List 2026-10-03.csv"
            path.write_text(content, encoding="utf-8")
            parsed = parse_price_list(path)
        self.assertEqual(parsed["status"], "completed")
        self.assertEqual(parsed["date_source"], "supplier_filename")
        self.assertEqual(parsed["products"][0]["normalized_identifier"], "012345678905")
        self.assertEqual(parsed["products"][0]["supplier_price"], "19.99")


if __name__ == "__main__":
    unittest.main()
