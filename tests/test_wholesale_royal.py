import contextlib
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "integrations"))
from import_royal_price_list import main
from wholesale_repository import WholesaleRepository
from wholesale_royal import availability, identifier_metadata, parse_workbook, product_identity
from wholesale_fixtures import GAME_A, GAME_B, royal_workbook


class RoyalParserTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "list.xlsx"

    def parse(self, rows, dated="2026-09-09"):
        return parse_workbook(royal_workbook(self.path, rows), dated)

    def test_workbook_raw_values_currency_summary_and_used(self):
        rows = [GAME_A, GAME_B,
                (" P4 Game \t used  ", "P4", "045496906023", 7, 25),
                ("Unused title is not USED marker", "PS5", "12345678901", 2, 1)]
        # Separate an ordinary title containing a substring from a standalone condition marker.
        rows[-1] = ("Unused title", "PS5", "12345678901", 2, 1)
        parsed = self.parse(rows)
        self.assertEqual(parsed["status"], "completed")
        self.assertEqual(parsed["summary"], {
            "rows_encountered": 4, "rows_imported": 3, "used_rows_skipped": 1,
            "non_product_rows_skipped": 2, "invalid_rows_skipped": 0, "duplicate_rows_skipped": 0,
            "conflict_rows_skipped": 0, "duplicate_price_superseded": 0,
            "identifier_warnings": 1, "availability_warnings": 0,
            "products_accepted": 3, "warnings": 1, "errors": 0,
        })
        a, b, questionable = parsed["products"]
        self.assertEqual(a["raw_identifier"], "045496906023")
        self.assertEqual(a["raw_title"], GAME_A[0])
        self.assertEqual((a["currency"], a["availability_raw"], a["availability_min"], a["availability_is_exact"]),
                         ("USD", "144+", 144, False))
        self.assertEqual((b["availability_min"], b["availability_is_exact"]), (74, True))
        self.assertEqual(b["normalized_identifier"], "710425597527")
        self.assertEqual(b["raw_identifier"], "7-10425-59752-7")
        self.assertEqual(questionable["normalized_identifier"], "12345678901")
        self.assertIsNone(questionable["check_digit_valid"])

    def test_check_digits_are_metadata_not_admission(self):
        for code, expected in (("045496906023", "valid"), ("4006381333931", "valid"),
                               ("045496906024", "invalid_check_digit"),
                               ("4006381333932", "invalid_check_digit"),
                               ("12345678901", "unexpected_length"), ("ABC-123", "non_barcode")):
            with self.subTest(code=code):
                p = self.parse([(" Game ", " SW ", code, 4, 0)])
                self.assertEqual(p["status"], "completed")
                self.assertEqual(p["products"][0]["identifier_status"], expected)
                self.assertEqual(p["products"][0]["raw_title"], " Game ")
        self.assertEqual(identifier_metadata(" 7-10425-59752-7 ")["normalized_identifier"], "710425597527")

    def test_numeric_identifier_rejected_instead_of_padding(self):
        p = self.parse([("Game", "SW", 45496906023, 4, 3)])
        self.assertEqual(p["status"], "rejected")
        self.assertEqual(p["summary"]["invalid_rows_skipped"], 1)
        self.assertEqual(p["source_rows"][0]["values"][2], "45496906023")
        self.assertEqual(p["source_rows"][0]["cell_types"][2], "n")

    def test_identity_is_supplier_offer_not_barcode_or_title_alone(self):
        base = product_identity("Game Deluxe", "PS5", "012345678905")
        self.assertEqual(base, product_identity(" game  deluxe ", "ps5 ", "012345678905"))
        for title, system, code in (("Game Standard", "PS5", "012345678905"),
                                    ("Game Deluxe", "P4", "012345678905"),
                                    ("Game Deluxe", "PS5", "012345678906")):
            self.assertNotEqual(base, product_identity(title, system, code))

    def test_exact_duplicates_keep_provenance_and_later_price_wins(self):
        p = self.parse([GAME_A, GAME_A])
        self.assertEqual(len(p["products"]), 1)
        self.assertEqual(p["products"][0]["source_row_numbers"], [2, 3])
        self.assertEqual(p["summary"]["duplicate_rows_skipped"], 1)
        conflict = self.parse([GAME_A, (*GAME_A[:3], 17, "144+")])
        self.assertEqual(conflict["status"], "completed")
        self.assertEqual(conflict["summary"]["conflict_rows_skipped"], 0)
        self.assertEqual(conflict["summary"]["duplicate_price_superseded"], 1)
        self.assertEqual(conflict["summary"]["rows_imported"], 1)
        self.assertEqual(conflict["products"][0]["supplier_price"], "17.00")
        self.assertEqual(conflict["products"][0]["source_row_numbers"], [2, 3])
        self.assertEqual([r["outcome"] for r in conflict["source_rows"][:2]], ["superseded", "accepted"])

        partial = self.parse([GAME_A, (*GAME_A[:3], 17, "144+"), GAME_B])
        self.assertEqual(partial["status"], "completed")
        self.assertEqual(partial["summary"]["rows_imported"], 2)
        self.assertEqual(partial["summary"]["conflict_rows_skipped"], 0)
        self.assertEqual(partial["warnings"][-1]["kind"], "duplicate_price_superseded")

    def test_explicit_and_document_dates_no_filename_or_mtime_inference(self):
        path = royal_workbook(self.path, [GAME_A])
        with self.assertRaisesRegex(ValueError, "effective-date"):
            parse_workbook(path)
        p = parse_workbook(path, "2026-09-09")
        self.assertEqual(p["date_source"], "operator_parameter")
        royal_workbook(path, [GAME_A], "09/09/2026")
        self.assertEqual(parse_workbook(path)["date_source"], "supplier_document")
        with self.assertRaisesRegex(ValueError, "conflicts"):
            parse_workbook(path, "2026-09-10")

    def test_availability_unknown_is_not_zero_or_exact(self):
        self.assertEqual(availability(" 144 + "), {"availability_min": 144, "availability_is_exact": False})
        self.assertEqual(availability("0"), {"availability_min": 0, "availability_is_exact": True})
        p = self.parse([(*GAME_A[:4], "CALL")])
        self.assertEqual(p["summary"]["availability_warnings"], 1)
        self.assertIsNone(p["products"][0]["availability_min"])

    def test_document_date_rows_are_metadata_not_products(self):
        path = royal_workbook(self.path, [GAME_A])
        book = load_workbook(path)
        book.active.append([None, "Price list date: 09/09/2026"])
        book.active.append([None, "2026-09-09"])
        book.defined_names.add(DefinedName("PRICE_LIST_DATE", attr_text="'COMPLETE LIST'!$B$6"))
        book.save(path)
        book.close()
        parsed = parse_workbook(path)
        self.assertEqual(parsed["status"], "completed")
        self.assertEqual(parsed["summary"]["rows_encountered"], 1)

    def test_order_and_sub_formulas_do_not_change_product_data(self):
        path = royal_workbook(self.path, [GAME_A])
        before = parse_workbook(path, "2026-09-09")
        book = load_workbook(path)
        book.active["A2"] = 99
        book.active["G2"] = "=E2*A2+123"
        book.save(path)
        book.close()
        after = parse_workbook(path, "2026-09-09")
        self.assertEqual(before["products"], after["products"])
        self.assertNotEqual(before["file_sha256"], after["file_sha256"])

    def test_supplier_quantity_total_formula_is_metadata(self):
        path = royal_workbook(self.path, [GAME_A])
        book = load_workbook(path)
        final_row = book.active.max_row + 1
        book.active.cell(final_row, 6, f"=SUM(F2:F{final_row - 1})")
        book.save(path)
        book.close()

        parsed = parse_workbook(path, "2026-09-09")

        self.assertEqual(parsed["status"], "completed")
        self.assertEqual(parsed["summary"]["rows_encountered"], 1)
        self.assertEqual(parsed["summary"]["non_product_rows_skipped"], 3)
        self.assertEqual(parsed["source_rows"][-1]["outcome"], "non_product")

    def test_price_formula_missing_identity_and_empty_list_fail_safely(self):
        for row in [(*GAME_A[:3], -1, 4), (*GAME_A[:3], "=1+2", 4),
                    (*GAME_A[:3], 1.234, 4), ("Game", "SW", "", 4, 3)]:
            self.assertEqual(self.parse([row])["status"], "rejected")
        self.assertEqual(self.parse([])["status"], "rejected")

    def test_preview_never_constructs_database_client(self):
        royal_workbook(self.path, [GAME_A])
        with patch("import_royal_price_list.get_repository") as repo, contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main([str(self.path), "--effective-date", "2026-09-09"]), 0)
            repo.assert_not_called()

    def test_repository_is_bounded_and_writes_through_rpc(self):
        client = MagicMock()
        repository = WholesaleRepository(client)
        repository.apply_import({"test": True})
        client.rpc.assert_called_once_with("wholesale_apply_import", {
            "p_payload": {"test": True}, "p_replaces_import_id": None})
        client.table.assert_not_called()
        with self.assertRaises(ValueError):
            repository.history("id", limit=501)
        with self.assertRaises(ValueError):
            repository.list_products("supplier", offset=-1)


if __name__ == "__main__":
    unittest.main()
